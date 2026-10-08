"""In-memory resume text extractor supporting PDF, DOCX, and TXT formats."""

import io
import logging
from pathlib import Path

import docx
import pypdf

logger = logging.getLogger(__name__)


class ResumeParseError(ValueError):
    """Exception raised when resume file cannot be parsed or is empty."""

    pass


def extract_text_from_pdf(pdf_bytes: bytes) -> str:
    """Extract plain text from PDF bytes in memory."""
    if not pdf_bytes:
        raise ResumeParseError("PDF file is empty.")

    try:
        reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
        pages_text: list[str] = []
        for _idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                pages_text.append(text.strip())

        extracted = "\n\n".join(pages_text).strip()
        if not extracted:
            raise ResumeParseError("Could not extract any readable text from PDF file.")
        return extracted
    except Exception as exc:
        if isinstance(exc, ResumeParseError):
            raise
        raise ResumeParseError(f"Failed to parse PDF document: {exc}") from exc


def extract_text_from_docx(docx_bytes: bytes) -> str:
    """Extract plain text from DOCX bytes in memory."""
    if not docx_bytes:
        raise ResumeParseError("DOCX file is empty.")

    try:
        doc = docx.Document(io.BytesIO(docx_bytes))
        paragraphs_text: list[str] = []

        for p in doc.paragraphs:
            if p.text.strip():
                paragraphs_text.append(p.text.strip())

        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    paragraphs_text.append(" | ".join(row_text))

        extracted = "\n".join(paragraphs_text).strip()
        if not extracted:
            raise ResumeParseError("Could not extract any readable text from DOCX file.")
        return extracted
    except Exception as exc:
        if isinstance(exc, ResumeParseError):
            raise
        raise ResumeParseError(f"Failed to parse DOCX document: {exc}") from exc


def extract_txt_text(txt_bytes: bytes) -> str:
    """Extract plain text from TXT bytes in memory."""
    if not txt_bytes:
        raise ResumeParseError("Text file is empty.")
    try:
        text = txt_bytes.decode("utf-8").strip()
        if not text:
            raise ResumeParseError("Text file is empty.")
        return text
    except UnicodeDecodeError as exc:
        raise ResumeParseError(f"Invalid text encoding in file: {exc}") from exc


extract_pdf_text = extract_text_from_pdf
extract_docx_text = extract_text_from_docx


def extract_resume_text(filename: str, file_bytes: bytes) -> str:
    """Extract plain text from uploaded file bytes strictly in memory."""
    ext = Path(filename).suffix.lower()

    if ext == ".pdf":
        return extract_text_from_pdf(file_bytes)
    elif ext in (".docx", ".doc"):
        return extract_text_from_docx(file_bytes)
    elif ext in (".txt", ".md"):
        return extract_txt_text(file_bytes)
    else:
        raise ResumeParseError(
            f"Unsupported file format '{ext}'. Please upload a PDF, DOCX, or TXT file."
        )
